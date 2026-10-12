"""
dataset_builder.py - Turn (almost) any pile of data into a clean, training-ready LLM dataset - at any scale.

SOURCES (mix freely, as many as you like):
  --source      files / folders / globs / .zip: .txt .md .rst .log .tex .srt, 40+ code languages, .ipynb, .pdf,
                .docx, .pptx, .epub, .html, .xml, .rtf, .csv, .tsv, .xlsx, .json, .jsonl, .parquet
                (+ images & scanned PDFs with --ocr, + audio/video with --transcribe)
  --url         web pages / linked files, optional same-site crawl (respects robots.txt)
  --hf          Hugging Face datasets, streamed: "owner/name" or "owner/name|subset|split"
  --github      repos: "owner/repo", "https://github.com/owner/repo/tree/branch" (downloads a zip, no git)
  --wikipedia   "Article title", "search: query", "category: Name", or a wikipedia.org URL
  --youtube     video / playlist / channel URLs -> transcripts (Whisper fallback with --transcribe)

PIPELINE (streaming, parallel, resumable):
  1. extract    -> build/extracted.jsonl   rows auto-classified as chat / instruction / preference / text
  2. qualify    -> build/kept.jsonl        unicode repair, cleanup, PII redaction, quality filters,
                                           exact + near-duplicate removal, token-aware chunking (multi-process)
  3. generate   (optional) local LLM writes Q&A pairs from raw documents
  4. finalize   -> train/eval as .jsonl or sharded .parquet, stats.json, README.md

  --resume continues an interrupted/stopped build: finished sources and phases are skipped.
  Disk-space and memory guards stop or degrade gracefully instead of crashing the machine.

TARGET FORMATS (--format): auto | pretrain {"text"} | sft {"messages"} | dpo {"prompt","chosen","rejected"}

Examples:
  python dataset_builder.py --name mydocs --source C:\\docs --source notes.zip --format pretrain
  python dataset_builder.py --name kb-qa --source C:\\manuals --ocr --format sft --synthetic
  python dataset_builder.py --name code --github huggingface/transformers --format pretrain --output_format parquet
  python dataset_builder.py --name wiki --wikipedia "search: large language models" --wiki_per_query 50
  python dataset_builder.py --name talks --youtube https://www.youtube.com/playlist?list=... --transcribe
  python dataset_builder.py --name big --hf HuggingFaceFW/fineweb-edu --max_rows 5000000 --output_format parquet --resume
"""

import os
import re
import sys
import json
import math
import time
import glob
import shutil
import random
import hashlib
import zipfile
import argparse
import tempfile
import traceback
import data_quality as dq
import multiprocessing as mp
from collections import Counter, deque
from urllib.parse import urljoin, urlparse
from concurrent.futures import ProcessPoolExecutor, as_completed

DATA_DIR = "data"

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".tex", ".org", ".adoc", ".srt", ".vtt"}
CODE_EXT = {".py", ".ipynb", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".go",
            ".rs", ".rb", ".php", ".sh", ".bash", ".ps1", ".bat", ".sql", ".yaml", ".yml", ".toml", ".ini", ".cfg",
            ".css", ".scss", ".swift", ".kt", ".scala", ".lua", ".r", ".m", ".pl", ".dart", ".vue", ".svelte",
            ".hs", ".ex", ".exs", ".clj", ".jl", ".zig", ".nim", ".f90", ".asm", ".gradle", ".cmake", ".dockerfile"}
TABLE_EXT = {".csv", ".tsv", ".xlsx", ".xls", ".json", ".jsonl", ".ndjson", ".parquet"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".gif"}
MEDIA_EXT = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".opus", ".aac", ".wma", ".mp4", ".mkv", ".mov", ".webm", ".avi", ".m4v"}
SKIP_EXT = {".ico", ".svg", ".exe", ".dll", ".so", ".bin", ".pt", ".pth", ".safetensors", ".gguf", ".onnx", ".pyc",
            ".o", ".obj", ".lib", ".class", ".jar", ".woff", ".woff2", ".ttf", ".otf", ".db", ".sqlite", ".lock",
            ".npy", ".npz", ".pkl", ".h5", ".ckpt", ".iso", ".dmg", ".msi", ".7z", ".rar", ".tar", ".gz"}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", ".idea", ".vscode", "dist", "build", ".next",
             "site-packages", ".mypy_cache", ".pytest_cache", "runs", "output_qlora"}

PROMPT_KEYS = ["instruction", "prompt", "question", "query", "input", "human", "user", "source", "problem", "request"]
CONTEXT_KEYS = ["context", "input", "passage", "document"]
RESPONSE_KEYS = ["output", "response", "answer", "completion", "assistant", "gpt", "target", "solution", "reply", "chosen_response"]
TEXT_KEYS = ["text", "content", "body", "article", "document", "passage", "markdown", "story", "code"]


# =============================================================================== status / progress
class Status:
    def __init__(self, build_dir, args, resume=False):
        self.path = os.path.join(build_dir, "status.json")
        self.t0 = time.time()
        self._last = 0
        prev = {}
        if resume and os.path.exists(self.path):
            with open(self.path, encoding="utf-8") as f:
                prev = json.load(f)
        self.s = {"pid": os.getpid(), "state": "running", "phase": "Starting", "started_at": self.t0,
                  "config": {k: v for k, v in vars(args).items() if isinstance(v, (str, int, float, bool, list, type(None)))},
                  "progress": {"done": 0, "total": 0, "unit": ""}, "counters": prev.get("counters", {}),
                  "dropped": prev.get("dropped", {}), "sources": prev.get("sources", {}), "warnings": [],
                  "resumed": bool(prev), "error": None}
        self.write(True)

    def phase(self, name, total=0, unit=""):
        print(f"\n==> {name}", flush=True)
        self.s["phase"] = name
        self.s["phase_started"] = time.time()
        self.s["progress"] = {"done": 0, "total": total, "unit": unit}
        self.write(True)

    def tick(self, n=1, total=None):
        p = self.s["progress"]
        p["done"] += n
        if total is not None:
            p["total"] = total
        self.write()

    def count(self, key, n=1):
        self.s["counters"][key] = self.s["counters"].get(key, 0) + n

    def drop(self, reason, n=1):
        self.s["dropped"][reason] = self.s["dropped"].get(reason, 0) + n

    def source(self, kind, n=1):
        self.s["sources"][kind] = self.s["sources"].get(kind, 0) + n

    def warn(self, msg):
        print(f"  ! {msg}", flush=True)
        if msg not in self.s["warnings"]:
            self.s["warnings"].append(msg)
        self.write(True)

    def write(self, force=False):
        now = time.time()
        if not force and now - self._last < 0.7:
            return
        self._last = now
        p = self.s["progress"]
        el = now - self.s.get("phase_started", now)
        p["eta_s"] = (el / p["done"]) * (p["total"] - p["done"]) if p["done"] and p["total"] and p["total"] > p["done"] else None
        self.s["elapsed_s"] = now - self.t0
        self.s["updated_at"] = now
        tmp = self.path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.s, f)
        for _ in range(5):
            try:
                os.replace(tmp, self.path)
                break
            except PermissionError:
                time.sleep(0.05)


class NullStatus(Status):
    """Status that only prints - used by library operations that report elsewhere."""
    def __init__(self):
        self.s = {"counters": {}, "dropped": {}, "sources": {}, "warnings": [], "progress": {"done": 0, "total": 0}}

    def write(self, force=False):
        pass


# =============================================================================== resource guards
def check_disk(path, min_free_gb):
    free = shutil.disk_usage(os.path.abspath(path)).free / 1024**3
    if free < min_free_gb:
        raise RuntimeError(f"Low disk space: only {free:.1f} GB free (minimum {min_free_gb} GB). "
                           "Free some space, then re-run with --resume to continue where it stopped.")
    return free


def ram_percent():
    try:
        import psutil
        return psutil.virtual_memory().percent
    except Exception:
        return 0


# =============================================================================== cleaning / quality
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_MULTI_NL = re.compile(r"\n{3,}")
_SPACES = re.compile(r"[ \t\u00a0]+")
PII = [
    (re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b"), "[EMAIL]"),
    (re.compile(r"\b(?:\d[ -]?){13,16}\b"), "[CARD]"),
    (re.compile(r"\b\d{3}-\d{2}-\d{4}\b"), "[SSN]"),
    (re.compile(r"(?<!\w)(?:\+?\d{1,3}[ .-]?)?\(?\d{3}\)?[ .-]?\d{3}[ .-]?\d{4}(?!\w)"), "[PHONE]"),
    (re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b"), "[IP]"),
    (re.compile(r"\b(?:sk|pk|ghp|gho|xox[bap])[-_][A-Za-z0-9_-]{16,}\b"), "[SECRET]"),
]


def clean_text(text, redact_pii=False, is_code=False):
    if not text:
        return ""
    try:
        import ftfy
        text = ftfy.fix_text(text)
    except Exception:
        pass
    text = _CTRL.sub("", text.replace("\r\n", "\n").replace("\r", "\n"))
    if not is_code:
        text = "\n".join(_SPACES.sub(" ", line).strip() for line in text.split("\n"))
    else:
        text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = _MULTI_NL.sub("\n\n", text).strip()
    if redact_pii:
        for rx, repl in PII:
            text = rx.sub(repl, text)
    return text


def quality_issue(text, min_chars, is_code=False):
    """Returns a drop reason, or None if the text is fine."""
    n = len(text)
    if n < min_chars:
        return "too short"
    if is_code:
        return None
    letters = sum(c.isalpha() for c in text)
    if letters / n < 0.4:
        return "mostly symbols/numbers"
    # Only substantive lines count - short structural lines ("|", "----", "}") repeat legitimately.
    lines = [l.strip() for l in text.split("\n") if len(l.strip()) >= 20]
    if len(lines) >= 5:
        dup_frac = 1 - len(set(lines)) / len(lines)
        if dup_frac > 0.5:
            return "repetitive lines"
    words = text.split()
    if words:
        if len(words) >= 50 and len(set(words)) / len(words) < 0.15:
            return "repetitive words"
        avg = sum(len(w) for w in words) / len(words)
        if avg > 15 or avg < 2:
            return "gibberish (word length)"
    low = text.lower()
    if n < 2000 and sum(k in low for k in ("cookie", "javascript", "sign in", "log in", "subscribe", "privacy policy")) >= 3:
        return "web boilerplate"
    return None


class Deduper:
    """Exact dedup (8-byte hashes) + MinHash near-dedup, with memory caps for very large builds."""

    def __init__(self, near=True, threshold=0.85, near_max=3_000_000, status=None):
        self.exact = set()
        self.lsh = None
        self.n = 0
        self.near_max = near_max
        self.status = status
        if near:
            try:
                from datasketch import MinHashLSH
                self.lsh = MinHashLSH(threshold=threshold, num_perm=64)
            except ImportError:
                pass

    def _disable_near(self, why):
        self.lsh = None
        import gc
        gc.collect()
        if self.status:
            self.status.warn(f"Near-duplicate detection turned off ({why}); exact dedup continues.")

    def is_dup(self, text):
        norm = re.sub(r"\W+", " ", text.lower()).strip()
        h = int.from_bytes(hashlib.blake2b(norm.encode(), digest_size=8).digest(), "little")
        if h in self.exact:
            return "exact duplicate"
        self.exact.add(h)
        if self.lsh is not None and len(norm) > 200:
            from datasketch import MinHash
            words = norm.split()
            m = MinHash(num_perm=64)
            m.update_batch([" ".join(words[i:i + 5]).encode() for i in range(max(1, len(words) - 4))])
            if self.lsh.query(m):
                return "near duplicate"
            self.n += 1
            self.lsh.insert(str(self.n), m)
            if self.n >= self.near_max:
                self._disable_near(f"over {self.near_max:,} documents")
            elif self.n % 20000 == 0 and ram_percent() > 90:
                self._disable_near("system RAM above 90%")
        return None


# =============================================================================== record classification
def _get(rec, keys):
    lower = {str(k).lower(): k for k in rec}
    for k in keys:
        if k in lower:
            v = rec[lower[k]]
            if isinstance(v, str) and v.strip():
                return v
    return None


def _as_text(v):
    if isinstance(v, str):
        return v
    if isinstance(v, list) and v and isinstance(v[-1], dict):  # chat-formatted chosen/rejected
        return v[-1].get("content", "")
    return "" if v is None else str(v)


def classify_record(rec):
    """Map an arbitrary dict to {"kind": "sft"|"dpo"|"text", ...}."""
    if not isinstance(rec, dict):
        return {"kind": "text", "text": str(rec)}
    keys = {str(k).lower() for k in rec}

    if {"chosen", "rejected"} <= keys:
        return {"kind": "dpo", "prompt": _as_text(rec.get("prompt") or rec.get("question") or ""),
                "chosen": _as_text(rec["chosen"]), "rejected": _as_text(rec["rejected"])}

    for key in ("messages", "conversations", "conversation", "dialog", "chat"):
        turns = rec.get(key)
        if isinstance(turns, list) and turns and isinstance(turns[0], dict):
            role_map = {"human": "user", "user": "user", "gpt": "assistant", "assistant": "assistant", "bot": "assistant",
                        "model": "assistant", "system": "system", "chatgpt": "assistant"}
            msgs = [{"role": role_map.get(str(t.get("role", t.get("from", "user"))).lower(), "user"),
                     "content": str(t.get("content", t.get("value", t.get("text", ""))))} for t in turns]
            return {"kind": "sft", "messages": msgs}

    prompt = _get(rec, PROMPT_KEYS)
    response = _get(rec, RESPONSE_KEYS)
    if prompt and response and prompt != response:
        ctx = _get(rec, CONTEXT_KEYS)
        if ctx and ctx not in (prompt, response):
            prompt = f"{prompt}\n\n{ctx}"
        msgs = []
        system = _get(rec, ["system", "system_prompt"])
        if system:
            msgs.append({"role": "system", "content": system})
        msgs += [{"role": "user", "content": prompt}, {"role": "assistant", "content": response}]
        return {"kind": "sft", "messages": msgs}

    text = _get(rec, TEXT_KEYS)
    if text:
        title = _get(rec, ["title", "name", "heading"])
        return {"kind": "text", "text": f"{title}\n\n{text}" if title and title not in text else text}

    parts = [f"{k}: {v}" for k, v in rec.items() if v not in (None, "") and not isinstance(v, (bytes, bytearray))]
    return {"kind": "text", "text": "\n".join(parts), "is_row": True}


# =============================================================================== extraction (worker processes)
def _read_text_file(path):
    with open(path, "rb") as f:
        raw = f.read()
    if b"\x00" in raw[:8192]:
        raise ValueError("binary file")
    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _html_to_text(html, url=None):
    try:
        import trafilatura
        txt = trafilatura.extract(html, url=url, include_tables=True, include_comments=False, favor_recall=True)
        if txt:
            return txt
    except Exception:
        pass
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "lxml")
    for t in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        t.decompose()
    return soup.get_text("\n")


def extract_document(path, ocr=False):
    """Extract text from one document-like file. Returns (docs, error)."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in IMAGE_EXT:
            from dataset_sources import ocr_image
            text = ocr_image(path)
            return ([{"kind": "text", "text": text, "src": path, "ocr": True}], None) if text.strip() else ([], "no text found in image")
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            pages = [(p.extract_text() or "") for p in reader.pages]
            text = "\n\n".join(pages)
            if len(text.strip()) < 50 * max(1, len(pages)) ** 0.5:
                if not ocr:
                    return [], "PDF has no text layer (scanned image - needs OCR)"
                from dataset_sources import ocr_pdf
                text = ocr_pdf(path)
                return [{"kind": "text", "text": text, "src": path, "ocr": True}], None
        elif ext == ".docx":
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                bits = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
                if slide.has_notes_slide:
                    bits.append(slide.notes_slide.notes_text_frame.text)
                slides.append(f"Slide {i}:\n" + "\n".join(b for b in bits if b.strip()))
            text = "\n\n".join(slides)
        elif ext == ".epub":
            import ebooklib
            from ebooklib import epub
            book = epub.read_epub(path)
            text = "\n\n".join(_html_to_text(item.get_content().decode("utf-8", "replace"))
                               for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT))
        elif ext in (".html", ".htm", ".xhtml", ".xml"):
            text = _html_to_text(_read_text_file(path))
        elif ext == ".rtf":
            text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", _read_text_file(path))
        elif ext == ".ipynb":
            nb = json.loads(_read_text_file(path))
            cells = []
            for c in nb.get("cells", []):
                src = "".join(c.get("source", []))
                cells.append(f"```python\n{src}\n```" if c.get("cell_type") == "code" else src)
            return [{"kind": "text", "text": "\n\n".join(cells), "is_code": True, "src": path}], None
        else:
            text = _read_text_file(path)
            is_code = ext in CODE_EXT or os.path.basename(path).lower() in ("dockerfile", "makefile")
            if is_code:
                text = f"File: {os.path.basename(path)}\n```{ext.lstrip('.')}\n{text}\n```"
            return [{"kind": "text", "text": text, "is_code": is_code, "src": path}], None
        return [{"kind": "text", "text": text, "src": path}], None
    except Exception as e:
        return [], f"{type(e).__name__}: {e}"


def iter_table_records(path, progress=None):
    """Stream rows from tabular / JSON files. progress(bytes) is called as data is consumed."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".jsonl", ".ndjson"):
        with open(path, "rb") as f:
            for raw in f:
                if progress:
                    progress(len(raw))
                line = raw.decode("utf-8", "replace").strip()
                if line:
                    try:
                        yield json.loads(line)
                    except json.JSONDecodeError:
                        yield {"text": line}
        return
    size = os.path.getsize(path)
    if ext == ".json":
        with open(path, encoding="utf-8", errors="replace") as f:
            data = json.load(f)
        if isinstance(data, dict):
            lists = [v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
            data = lists[0] if lists else [data]
        yield from (data if isinstance(data, list) else [data])
    elif ext in (".csv", ".tsv"):
        import pandas as pd
        for chunk in pd.read_csv(path, sep="\t" if ext == ".tsv" else None, engine="python", chunksize=5000,
                                 dtype=str, keep_default_na=False, on_bad_lines="skip", encoding_errors="replace"):
            yield from chunk.to_dict("records")
    elif ext in (".xlsx", ".xls"):
        import pandas as pd
        for _, df in pd.read_excel(path, sheet_name=None, dtype=str).items():
            yield from df.fillna("").to_dict("records")
    elif ext == ".parquet":
        import pyarrow.parquet as pq
        pf = pq.ParquetFile(path)
        per_group = size / max(1, pf.num_row_groups)
        for i in range(pf.num_row_groups):
            for batch in pf.read_row_group(i).to_batches(max_chunksize=2000):
                yield from batch.to_pylist()
            if progress:
                progress(per_group)
        return
    if progress:
        progress(size)


# =============================================================================== discovery
def discover_files(sources, tmp_root, status, ocr=False, transcribe=False):
    files = []
    for src in sources:
        src = src.strip().strip('"')
        if not src:
            continue
        matches = glob.glob(src, recursive=True) if any(c in src for c in "*?[") else [src]
        for m in matches:
            if os.path.isdir(m):
                for root, dirs, names in os.walk(m):
                    dirs[:] = [d for d in dirs if d not in SKIP_DIRS and not d.startswith(".")]
                    files += [os.path.join(root, n) for n in names]
            elif os.path.isfile(m):
                files.append(m)
            else:
                status.warn(f"source not found: {m}")
                status.drop("source not found")
    out = []
    for f in files:
        ext = os.path.splitext(f)[1].lower()
        if ext == ".zip":
            dest = os.path.join(tmp_root, "unzipped", hashlib.md5(os.path.abspath(f).encode()).hexdigest()[:10])
            try:
                with zipfile.ZipFile(f) as z:
                    for member in z.namelist():  # zip-slip guard
                        target = os.path.realpath(os.path.join(dest, member))
                        if not target.startswith(os.path.realpath(dest)):
                            raise ValueError(f"unsafe path in zip: {member}")
                    z.extractall(dest)
                out += discover_files([dest], tmp_root, status, ocr, transcribe)
                status.source("zip archive")
            except Exception as e:
                status.warn(f"bad zip {os.path.basename(f)}: {e}")
                status.drop("unreadable zip")
        elif ext in SKIP_EXT or (ext in IMAGE_EXT and not ocr) or (ext in MEDIA_EXT and not transcribe):
            hint = " - enable OCR" if ext in IMAGE_EXT else " - enable transcription" if ext in MEDIA_EXT else ""
            status.drop(f"unsupported type ({ext}{hint})")
        elif os.path.getsize(f) == 0:
            status.drop("empty file")
        else:
            out.append(f)
    return out


# =============================================================================== web
_TRACKING = re.compile(r"^(utm_.*|fbclid|gclid|mc_cid|mc_eid|ref|ref_src|igshid|_hsenc|_hsmi)$", re.I)
_JUNK_PATH = re.compile(r"/(login|log-in|signin|sign-in|signup|register|cart|checkout|account|my-account|cdn-cgi|wp-json|wp-admin|feed|rss|"
                        r"print|share|calendar|tag|tags|author|search|comments?)(/|$|\?)", re.I)


def canonical_url(url):
    """Normalise so the same page isn't fetched twice (fragments, tracking params, trailing slash, host case)."""
    from urllib.parse import parse_qsl, urlencode, urlunparse
    u = urlparse(url.strip())
    q = sorted((k, v) for k, v in parse_qsl(u.query) if not _TRACKING.match(k))
    path = re.sub(r"/{2,}", "/", u.path or "/")
    if len(path) > 1 and path.endswith("/"):
        path = path[:-1]
    return urlunparse((u.scheme.lower(), u.netloc.lower(), path, "", urlencode(q), ""))


def fetch_urls(urls, depth, max_pages, same_site, status, tmp_root, ocr=False, delay=0.5, sitemap=False, workers=8):
    """Parallel BFS crawl: polite per-host rate limit, retries with backoff, canonical-URL dedup, junk-page skipping,
    robots.txt, response size cap, auto-sitemap discovery. Yields docs."""
    import requests
    import threading
    from urllib import robotparser
    from concurrent.futures import ThreadPoolExecutor
    UA = "TrainerBoi-DatasetBuilder/1.0 (+local research use)"
    local = threading.local()
    robots, host_next, lock = {}, {}, threading.Lock()
    MAX_BYTES = 25 * 1024 * 1024

    def session():
        if not hasattr(local, "s"):
            local.s = requests.Session()
            local.s.headers["User-Agent"] = UA
            local.s.headers["Accept-Language"] = "en,*;q=0.5"
        return local.s

    expanded = []
    for u in urls:
        u = u.strip()
        if not u:
            continue
        if sitemap or u.endswith(".xml"):
            try:
                from dataset_sources import sitemap_urls
                sm_urls = sitemap_urls(u, max_urls=max_pages)
                if sm_urls:
                    status.phase(f"Discovered {len(sm_urls)} URLs via sitemap for {u}")
                    expanded.extend(sm_urls)
                    continue
            except Exception:
                pass
        expanded.append(u)

    def allowed(url):
        pu = urlparse(url)
        base = f"{pu.scheme}://{pu.netloc}"
        with lock:
            if base not in robots:
                rp = robotparser.RobotFileParser(base + "/robots.txt")
                try:
                    rp.read()
                except Exception:
                    rp = None
                robots[base] = rp
            rp = robots[base]
        return rp is None or rp.can_fetch(UA, url)

    def polite_wait(host):
        with lock:
            now = time.time()
            at = max(now, host_next.get(host, 0))
            host_next[host] = at + delay
        if at > now:
            time.sleep(at - now)

    def fetch(url):
        """-> ("ok", response) | ("skip", reason) | ("fail", message)"""
        if not allowed(url):
            return "skip", "blocked by robots.txt"
        host = urlparse(url).netloc
        last = ""
        for attempt in range(3):
            polite_wait(host)
            try:
                r = session().get(url, timeout=(8, 25), stream=True)
                if r.status_code in (429, 500, 502, 503, 504):
                    last = f"HTTP {r.status_code}"
                    wait = r.headers.get("Retry-After", "")
                    time.sleep(min(30, float(wait)) if wait.isdigit() else 1.5 * (attempt + 1) ** 2)
                    continue
                r.raise_for_status()
                if int(r.headers.get("content-length") or 0) > MAX_BYTES:
                    return "skip", "file too large"
                body = r.raw.read(MAX_BYTES + 1, decode_content=True)
                if len(body) > MAX_BYTES:
                    return "skip", "file too large"
                r._content, r._content_consumed = body, True
                return "ok", r
            except Exception as e:
                last = str(e)[:160]
                time.sleep(1.0 * (attempt + 1))
        return "fail", last

    seen, pages = set(), 0
    frontier = []                       # (url, depth), breadth-first
    for u in expanded:
        cu = canonical_url(u)
        if cu not in seen:
            seen.add(cu)
            frontier.append((cu, 0))
    roots = {urlparse(u).netloc for u, _ in frontier}
    with ThreadPoolExecutor(max_workers=workers) as pool:
        while frontier and pages < max_pages:
            k = max(1, min(max_pages - pages, workers * 4))
            batch, frontier = frontier[:k], frontier[k:]
            for (url, d), (verdict, r) in zip(batch, pool.map(fetch, [u for u, _ in batch])):
                if pages >= max_pages:
                    break
                if verdict == "skip":
                    status.drop(r)
                    continue
                if verdict == "fail":
                    status.warn(f"{url}: {r}")
                    status.drop("URL fetch failed")
                    continue
                pages += 1
                status.tick(1, total=max(max_pages if depth else len(expanded), pages))
                ctype = r.headers.get("content-type", "").split(";")[0].strip()
                if "html" in ctype:
                    html = r.text
                    canon = re.search(r"<link[^>]+rel=[\"']canonical[\"'][^>]*href=[\"']([^\"']+)", html, re.I)
                    if canon:
                        cu = canonical_url(urljoin(url, canon.group(1)))
                        if cu != url and cu in seen:
                            status.drop("duplicate page (canonical)")
                            continue
                        seen.add(cu)
                    text = _html_to_text(html, url=url)
                    title = re.search(r"<title[^>]*>(.*?)</title>", html, re.S | re.I)
                    if title:
                        text = f"{title.group(1).strip()}\n\n{text}"
                    status.source("web page")
                    yield {"kind": "text", "text": text, "src": url}
                    if d < depth:
                        from bs4 import BeautifulSoup
                        for a in BeautifulSoup(html, "lxml").find_all("a", href=True):
                            nxt = urljoin(url, a["href"])
                            if not nxt.startswith("http"):
                                continue
                            nxt = canonical_url(nxt)
                            if nxt in seen or (same_site and urlparse(nxt).netloc not in roots):
                                continue
                            if _JUNK_PATH.search(urlparse(nxt).path + "/"):
                                continue
                            seen.add(nxt)
                            frontier.append((nxt, d + 1))
                else:
                    ext = os.path.splitext(urlparse(url).path)[1].lower() or {"application/pdf": ".pdf", "text/plain": ".txt",
                                                                              "application/json": ".json", "text/csv": ".csv"}.get(ctype, "")
                    if not ext or ext in SKIP_EXT or ext in MEDIA_EXT:
                        status.drop(f"unsupported URL type ({ctype})")
                        continue
                    dl_dir = os.path.join(tmp_root, "downloads")
                    os.makedirs(dl_dir, exist_ok=True)
                    fp = os.path.join(dl_dir, hashlib.md5(url.encode()).hexdigest()[:12] + ext)
                    with open(fp, "wb") as f:
                        f.write(r.content)
                    status.source(f"downloaded {ext}")
                    if ext in TABLE_EXT:
                        for rec in iter_table_records(fp):
                            yield {**classify_record(rec), "src": url}
                    else:
                        docs, err = extract_document(fp, ocr)
                        if err:
                            status.drop("extraction failed")
                        for doc in docs:
                            doc["src"] = url
                            yield doc

# =============================================================================== chunking
class Chunker:
    def __init__(self, tokenizer_id, chunk_tokens, overlap_tokens):
        from transformers import AutoTokenizer
        self.tok = AutoTokenizer.from_pretrained(tokenizer_id)
        self.size = chunk_tokens
        self.overlap = overlap_tokens

    def count(self, text):
        return len(self.tok(text, add_special_tokens=False)["input_ids"])

    def chunks(self, text):
        # Fast path: a token is at least ~1 char, so short texts can't exceed the budget.
        if len(text) <= self.size or self.count(text) <= self.size:
            return [text]
        units = []
        for para in re.split(r"\n\s*\n", text):
            if self.count(para) <= self.size:
                units.append(para)
            else:
                for sent in re.split(r"(?<=[.!?])\s+|\n", para):
                    if self.count(sent) <= self.size:
                        units.append(sent)
                    else:  # hard split pathological runs (minified code, tables)
                        ids = self.tok(sent, add_special_tokens=False)["input_ids"]
                        units += [self.tok.decode(ids[i:i + self.size]) for i in range(0, len(ids), self.size)]
        lens = [len(x) for x in self.tok(units, add_special_tokens=False)["input_ids"]] if units else []
        out, cur, cur_lens = [], [], []
        for u, n in zip(units, lens):
            if cur and sum(cur_lens) + n > self.size:
                out.append("\n\n".join(cur))
                # Carry the tail of the previous chunk forward as overlap (context continuity).
                keep, keep_lens = [], []
                for pu, pn in zip(reversed(cur), reversed(cur_lens)):
                    if sum(keep_lens) + pn > self.overlap:
                        break
                    keep.insert(0, pu)
                    keep_lens.insert(0, pn)
                cur, cur_lens = keep, keep_lens
            cur.append(u)
            cur_lens.append(n)
        if cur:
            out.append("\n\n".join(cur))
        return out


# =============================================================================== qualify (worker processes)
_Q = {}


def _qinit(cfg):
    _Q.clear()
    _Q.update(cfg)
    if cfg.get("chunk_tokens"):
        _Q["chunker"] = Chunker(cfg["tokenizer"], cfg["chunk_tokens"], cfg["overlap_tokens"])


def _auto_drop(example):
    """Shared quality rules (data_quality.py): returns a drop reason for flags that are safe to drop automatically."""
    import data_quality as dq
    _, flags = dq.check(example)
    hit = [f for f in flags if dq.FLAGS[f][2]]
    return dq.FLAGS[hit[0]][1] if hit else None


def qualify_doc(line):
    """One extracted doc -> list of ("keep", example, dedup_text) | ("synth", chunk, dedup_text) | ("drop", reason)."""
    c = _Q
    doc = json.loads(line)
    kind, fmt = doc["kind"], c["fmt"]
    out = []

    if kind == "dpo":
        if fmt != "dpo":
            return [("drop", "preference row (not used in this format)")]
        p, ch, rj = (clean_text(doc[k], c["redact_pii"]) for k in ("prompt", "chosen", "rejected"))
        if not ch or not rj or ch == rj:
            return [("drop", "invalid preference pair")]
        bad = _auto_drop({"prompt": p, "chosen": ch, "rejected": rj})
        if bad:
            return [("drop", bad)]
        return [("keep", {"prompt": p, "chosen": ch, "rejected": rj}, p + "\n" + ch)]

    if kind == "sft":
        msgs = [{"role": m["role"], "content": clean_text(m["content"], c["redact_pii"])} for m in doc["messages"]]
        msgs = [m for m in msgs if m["content"]]
        if fmt == "pretrain":
            doc = {"kind": "text", "text": "\n\n".join(f"{m['role'].title()}: {m['content']}" for m in msgs)}
            kind = "text"
        elif fmt == "sft":
            if not any(m["role"] == "assistant" for m in msgs) or not any(m["role"] == "user" for m in msgs):
                return [("drop", "conversation missing user or assistant turn")]
            reply = next(m["content"] for m in reversed(msgs) if m["role"] == "assistant")
            if len(reply) < c["min_response_chars"]:
                return [("drop", "response too short")]
            bad = _auto_drop({"messages": msgs})
            if bad:
                return [("drop", bad)]
            if c["system_prompt"] and msgs[0]["role"] != "system":
                msgs.insert(0, {"role": "system", "content": c["system_prompt"]})
            return [("keep", {"messages": msgs}, " ".join(m["content"] for m in msgs))]
        else:
            return [("drop", "instruction row (not used in this format)")]

    # raw text
    if fmt == "dpo":
        return [("drop", "raw text (not used for DPO)")]
    if fmt == "sft" and not c["synthetic"]:
        return [("drop", "raw text (enable synthetic Q&A to use)")]
    is_code = doc.get("is_code", False)
    text = clean_text(doc["text"], c["redact_pii"], is_code=is_code)
    reason = quality_issue(text, c["min_chars"], is_code)
    if reason:
        return [("drop", reason)]
    if not is_code:
        bad = _auto_drop({"text": text})
        if bad:
            return [("drop", bad)]
    for chunk in c["chunker"].chunks(text):
        reason = quality_issue(chunk, c["min_chars"], is_code)
        if reason:
            out.append(("drop", f"chunk: {reason}"))
        elif fmt == "pretrain":
            out.append(("keep", {"text": chunk}, chunk))
        else:
            out.append(("synth", chunk, chunk))
    return out


# =============================================================================== synthetic Q&A generation
QA_SYSTEM = "You are an expert at creating high-quality training data for AI assistants."
QA_PROMPT = """Read the passage below and write {n} diverse, high-quality question-answer pairs based ONLY on its content.

Rules:
- Questions must be self-contained: never say "the passage", "the text", "the author" or "according to".
- Mix question types: factual, explanatory ("why"/"how"), and practical/applied.
- Answers must be accurate, complete, and helpful (2-6 sentences), written as a knowledgeable assistant would answer.
- Skip trivia about formatting, page numbers, or navigation text.

Return ONLY a JSON array like: [{{"question": "...", "answer": "..."}}]

Passage:
\"\"\"
{chunk}
\"\"\""""


class QAGenerator:
    def __init__(self, model_id, load_4bit=False, batch_size=8, max_new_tokens=700):
        import torch
        from transformers import AutoTokenizer, AutoModelForCausalLM, BitsAndBytesConfig
        self.torch = torch
        self.tok = AutoTokenizer.from_pretrained(model_id, padding_side="left")
        if self.tok.pad_token is None:
            self.tok.pad_token = self.tok.eos_token
        quant = BitsAndBytesConfig(load_in_4bit=True, bnb_4bit_quant_type="nf4",
                                   bnb_4bit_compute_dtype=torch.bfloat16) if load_4bit else None
        self.model = AutoModelForCausalLM.from_pretrained(model_id, dtype=torch.bfloat16, quantization_config=quant,
                                                          device_map={"": 0}).eval()
        self.batch_size = batch_size
        self.max_new_tokens = max_new_tokens

    def generate(self, chunks, n):
        prompts = [self.tok.apply_chat_template(
            [{"role": "system", "content": QA_SYSTEM}, {"role": "user", "content": QA_PROMPT.format(n=n, chunk=c)}],
            tokenize=False, add_generation_prompt=True) for c in chunks]
        enc = self.tok(prompts, return_tensors="pt", padding=True, add_special_tokens=False).to(self.model.device)
        with self.torch.no_grad():
            out = self.model.generate(**enc, max_new_tokens=self.max_new_tokens, do_sample=True, temperature=0.7,
                                      top_p=0.9, pad_token_id=self.tok.pad_token_id)
        return self.tok.batch_decode(out[:, enc["input_ids"].shape[1]:], skip_special_tokens=True)


def parse_qa(text):
    pairs = []
    m = re.search(r"\[.*\]", text, re.S)
    if m:
        try:
            for item in json.loads(m.group(0)):
                if isinstance(item, dict):
                    pairs.append((str(item.get("question", "")).strip(), str(item.get("answer", "")).strip()))
        except json.JSONDecodeError:
            pass
    if not pairs:
        for q, a in re.findall(r"(?:Q|Question)\s*\d*[:.]\s*(.+?)\n+\s*(?:A|Answer)\s*\d*[:.]\s*(.+?)(?=\n\s*(?:Q|Question)\s*\d*[:.]|\Z)", text, re.S):
            pairs.append((q.strip(), a.strip()))
    return pairs


BAD_Q = re.compile(r"\b(the|this) (passage|text|document|excerpt|article|author)\b|according to", re.I)


# =============================================================================== finalize (shared with dataset_store merge)
class _JsonlWriter:
    def __init__(self, path):
        self.path = path
        self.f = open(path, "w", encoding="utf-8")
        self.files = [path]

    def write(self, line, ex):
        self.f.write(line if line.endswith("\n") else line + "\n")

    def close(self):
        self.f.close()


class _ParquetWriter:
    def __init__(self, out_dir, prefix, shard_rows):
        self.out_dir, self.prefix, self.shard_rows = out_dir, prefix, shard_rows
        self.rows, self.files = [], []

    def write(self, line, ex):
        self.rows.append(ex)
        if len(self.rows) >= self.shard_rows:
            self._flush()

    def _flush(self):
        if not self.rows:
            return
        import pyarrow as pa
        import pyarrow.parquet as pq
        path = os.path.join(self.out_dir, f"{self.prefix}-{len(self.files):05d}.parquet")
        # Small row groups = fast random access when browsing in the Dataset Library.
        pq.write_table(pa.Table.from_pylist(self.rows), path, compression="zstd", row_group_size=10_000)
        self.files.append(path)
        self.rows = []

    def close(self):
        self._flush()


def _example_text(e):
    if "text" in e:
        return e["text"]
    if "messages" in e:
        return "\n".join(m["content"] for m in e["messages"])
    return f"{e.get('prompt', '')}\n{e.get('chosen', '')}"


def finalize(kept_path, out_dir, build_dir, st, fmt, *, name, tokenizer_id, eval_ratio=0.02, seed=42,
             output_format="jsonl", shard_rows=500_000, min_free_gb=2.0, token_sample=200_000, extra=None):
    """Shuffle (bucketed, out-of-core), split, write jsonl or parquet shards, measure tokens, write stats + card."""
    rng = random.Random(seed)
    for old in glob.glob(os.path.join(out_dir, "train*")) + glob.glob(os.path.join(out_dir, "eval*")):
        os.remove(old)
    size = os.path.getsize(kept_path)
    total_lines = sum(1 for _ in open(kept_path, "rb"))
    st.phase("Shuffling, splitting & writing", total=total_lines, unit="examples")
    check_disk(out_dir, min_free_gb + size / 1024**3 * 2)

    n_buckets = max(1, math.ceil(size / (256 * 1024**2)))  # ~256 MB per in-memory bucket
    bucket_paths = [os.path.join(build_dir, f"bucket{b}.jsonl") for b in range(n_buckets)]
    bfs = [open(p, "w", encoding="utf-8") for p in bucket_paths]
    with open(kept_path, encoding="utf-8") as kf:
        for line in kf:
            bfs[rng.randrange(n_buckets)].write(line)
    for f in bfs:
        f.close()

    from transformers import AutoTokenizer
    tok = AutoTokenizer.from_pretrained(tokenizer_id)
    if output_format == "parquet":
        tw, vw = _ParquetWriter(out_dir, "train", shard_rows), _ParquetWriter(out_dir, "eval", shard_rows)
    else:
        tw, vw = _JsonlWriter(os.path.join(out_dir, "train.jsonl")), _JsonlWriter(os.path.join(out_dir, "eval.jsonl"))

    lengths, sample_chars, total_chars, n_train, n_eval, preview = [], 0, 0, 0, 0, []
    for bp in bucket_paths:
        with open(bp, encoding="utf-8") as f:
            lines = f.readlines()
        rng.shuffle(lines)
        for start in range(0, len(lines), 512):
            block = lines[start:start + 512]
            exs = [json.loads(l) for l in block]
            texts = [_example_text(e) for e in exs]
            total_chars += sum(len(t) for t in texts)
            if len(lengths) < token_sample:  # tokenize a large random sample; extrapolate the rest
                lengths += [len(x) for x in tok(texts, add_special_tokens=False)["input_ids"]]
                sample_chars += sum(len(t) for t in texts)
            for ex, line in zip(exs, block):
                if rng.random() < eval_ratio:
                    vw.write(line, ex)
                    n_eval += 1
                else:
                    tw.write(line, ex)
                    n_train += 1
                    if len(preview) < 30:
                        preview.append(ex)
            st.tick(len(block))
        os.remove(bp)
    tw.close()
    vw.close()

    estimated = n_train + n_eval > len(lengths)
    total_tokens = int(sum(lengths) / max(sample_chars, 1) * total_chars) if estimated else sum(lengths)
    lengths.sort()
    pct = lambda q: lengths[min(len(lengths) - 1, int(q * len(lengths)))]
    train_path = os.path.join(out_dir, "train.jsonl") if output_format == "jsonl" else out_dir
    stats = {
        "name": name, "format": fmt, "train_examples": n_train, "eval_examples": n_eval,
        "total_tokens": total_tokens, "tokens_estimated": estimated,
        "tokens_mean": sum(lengths) / len(lengths), "tokens_p50": pct(0.5), "tokens_p95": pct(0.95),
        "tokens_max": lengths[-1], "tokenizer": tokenizer_id, "output_format": output_format,
        "files": [os.path.relpath(p, out_dir).replace("\\", "/") for p in tw.files + vw.files],
        "length_histogram": _histogram(lengths), "built_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "train_path": train_path.replace("\\", "/"),
        "dropped": st.s["dropped"], "sources": st.s["sources"], **(extra or {}),
    }
    with open(os.path.join(out_dir, "stats.json"), "w", encoding="utf-8") as f:
        json.dump(stats, f, indent=2)
    with open(os.path.join(out_dir, "preview.json"), "w", encoding="utf-8") as f:
        json.dump(preview, f, indent=2, ensure_ascii=False)
    _write_card(out_dir, stats)
    return stats


def _histogram(lengths, bins=20):
    if not lengths:
        return []
    hi = lengths[min(len(lengths) - 1, int(0.99 * len(lengths)))] or 1
    width = max(1, -(-hi // bins))
    counts = Counter(min(n // width, bins - 1) for n in lengths)
    return [{"tokens": f"{i * width}-{(i + 1) * width}", "examples": counts.get(i, 0)} for i in range(bins)]


def _write_card(out_dir, s):
    dropped = "\n".join(f"| {k} | {v:,} |" for k, v in sorted(s.get("dropped", {}).items(), key=lambda x: -x[1])) or "| (none) | 0 |"
    est = " (estimated from a sample)" if s.get("tokens_estimated") else ""
    with open(os.path.join(out_dir, "README.md"), "w", encoding="utf-8") as f:
        f.write(f"""# Dataset: {s['name']}

Built {s['built_at']} by TrainerBoi

| | |
|---|---|
| Format | **{s['format']}** ({s['output_format']}) |
| Train examples | {s['train_examples']:,} |
| Eval examples | {s['eval_examples']:,} |
| Total tokens | {s['total_tokens']:,}{est} ({s['tokenizer']}) |
| Tokens / example | mean {s['tokens_mean']:.0f} - median {s['tokens_p50']} - p95 {s['tokens_p95']} - max {s['tokens_max']} |

## Sources
{chr(10).join(f"- {k}: {v:,}" for k, v in s.get('sources', {}).items())}

## Dropped during qualification
| Reason | Count |
|---|---|
{dropped}
""")


# =============================================================================== main build
def build(args):
    out_dir = os.path.join(DATA_DIR, args.name)
    build_dir = os.path.join(out_dir, "build")
    if os.path.exists(os.path.join(out_dir, "stats.json")) and not (args.resume or args.overwrite):
        raise SystemExit(f"Dataset '{args.name}' already exists. Use --overwrite to rebuild it, --resume to continue, "
                         "or pick another --name.")
    status_path = os.path.join(build_dir, "status.json")
    if args.resume and os.path.exists(status_path):
        # Resuming without re-typing sources: reuse the original build's settings.
        with open(status_path, encoding="utf-8") as f:
            prev = json.load(f).get("config", {})
        if not (args.source or args.url or args.hf or args.github or args.wikipedia or args.youtube):
            for k, v in prev.items():
                if k not in ("resume", "overwrite") and hasattr(args, k):
                    setattr(args, k, v)
        print(f"==> Resuming build '{args.name}'", flush=True)
    elif not args.resume and os.path.isdir(out_dir):
        _clean_dir(out_dir)
    os.makedirs(build_dir, exist_ok=True)
    st = Status(build_dir, args, resume=args.resume)
    try:
        _build(args, out_dir, build_dir, st)
    except KeyboardInterrupt:
        st.s.update(state="stopped", phase="Stopped")
        st.write(True)
        raise
    except Exception as e:
        st.s.update(state="error", phase="Crashed", error=f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}")
        st.write(True)
        raise


def _build(args, out_dir, build_dir, st):
    rng = random.Random(args.seed)
    tmp_root = os.path.join(build_dir, "tmp")
    os.makedirs(tmp_root, exist_ok=True)
    extracted_path = os.path.join(build_dir, "extracted.jsonl")
    done_path = os.path.join(build_dir, "sources_done.txt")
    phase_done = lambda p: os.path.exists(os.path.join(build_dir, f"phase_{p}.done"))
    mark_done = lambda p: open(os.path.join(build_dir, f"phase_{p}.done"), "w").close()

    done = set()
    if args.resume and os.path.exists(done_path):
        done = set(open(done_path, encoding="utf-8").read().splitlines())
        print(f"  resuming: {len(done):,} sources already processed", flush=True)
    done_f = open(done_path, "a", encoding="utf-8")

    def finish_source(key):
        done.add(key)
        done_f.write(key + "\n")
        done_f.flush()

    kinds = Counter()

    def emit(f, doc):
        kinds[doc["kind"]] += 1
        st.count(f"extracted: {doc['kind']}")
        f.write(json.dumps(doc, ensure_ascii=False) + "\n")

    # ================================================================== 1. extract
    if not phase_done("extract"):
        check_disk(out_dir, args.min_free_gb)
        st.phase("Discovering sources")
        files = discover_files(args.source or [], tmp_root, st, args.ocr, args.transcribe)

        if args.github:
            for spec in args.github:
                if f"gh:{spec}" in done:
                    continue
                try:
                    from dataset_sources import github_zip
                    st.phase(f"Downloading GitHub repo {spec}")
                    zp, label = github_zip(spec, os.path.join(tmp_root, "github"), args.github_max_mb)
                    repo_files = discover_files([zp], tmp_root, st, args.ocr, args.transcribe)
                    files += repo_files
                    st.source("GitHub repo")
                    print(f"  {label}: {len(repo_files):,} files", flush=True)
                except Exception as e:
                    st.warn(f"GitHub {spec}: {e}")
                    st.drop("GitHub repo failed")

        media = [f for f in files if os.path.splitext(f)[1].lower() in MEDIA_EXT]
        tables = [f for f in files if os.path.splitext(f)[1].lower() in TABLE_EXT]
        docs = [f for f in files if f not in set(media) | set(tables)]
        todo_docs = [f for f in docs if f"file:{os.path.abspath(f)}" not in done]
        print(f"  {len(docs):,} documents ({len(docs) - len(todo_docs):,} already done), {len(tables):,} tables, "
              f"{len(media):,} audio/video, {len(args.url or [])} URLs, {len(args.hf or [])} HF datasets", flush=True)
        st.s["counters"]["files found"] = len(files)

        with open(extracted_path, "a" if args.resume else "w", encoding="utf-8") as ef:
            # ---------------- documents (parallel)
            if todo_docs:
                st.phase("Extracting documents" + (" (with OCR)" if args.ocr else ""), total=len(todo_docs), unit="files")
                with ProcessPoolExecutor(max_workers=max(1, min(args.workers, len(todo_docs)))) as pool:
                    pending = set()
                    it = iter(todo_docs)
                    # Bounded submission keeps memory flat even with millions of files.
                    for f in it:
                        pending.add(pool.submit(_extract_job, f, args.ocr))
                        if len(pending) >= args.workers * 8:
                            break
                    while pending:
                        fut = next(as_completed(pending))
                        pending.remove(fut)
                        path, docs_out, err = fut.result()
                        if err:
                            st.drop("no text layer (needs OCR)" if "OCR" in err else
                                    "binary/unreadable file" if "binary" in err else "extraction failed")
                            if "OCR" not in err and "binary" not in err:
                                print(f"  ! {os.path.basename(path)}: {err}", flush=True)
                        ext = os.path.splitext(path)[1].lower() or "(none)"
                        st.source(f"file {ext}")
                        for d in docs_out:
                            emit(ef, d)
                        finish_source(f"file:{os.path.abspath(path)}")
                        st.tick()
                        if st.s["progress"]["done"] % 2000 == 0:
                            check_disk(out_dir, args.min_free_gb)
                            ef.flush()
                        nxt = next(it, None)
                        if nxt:
                            pending.add(pool.submit(_extract_job, nxt, args.ocr))

            # ---------------- tables / JSON (streamed, byte-level progress)
            todo_tables = [t for t in tables if f"table:{os.path.abspath(t)}" not in done]
            if todo_tables:
                total_mb = sum(os.path.getsize(t) for t in todo_tables) / 1e6
                st.phase("Reading tables & JSON", total=round(total_mb, 1), unit="MB")
                for path in todo_tables:
                    n = 0
                    try:
                        for rec in iter_table_records(path, progress=lambda b: st.tick(b / 1e6)):
                            doc = classify_record(rec)
                            doc["src"] = path
                            emit(ef, doc)
                            n += 1
                            if n % 50000 == 0:
                                check_disk(out_dir, args.min_free_gb)
                            if args.max_rows and n >= args.max_rows:
                                break
                    except Exception as e:
                        st.warn(f"{os.path.basename(path)}: {e}")
                        st.drop("unreadable table/json")
                    st.source(f"file {os.path.splitext(path)[1].lower()}")
                    print(f"  {os.path.basename(path)}: {n:,} rows", flush=True)
                    finish_source(f"table:{os.path.abspath(path)}")

            # ---------------- Hugging Face (streamed)
            for spec in args.hf or []:
                if f"hf:{spec}" in done:
                    continue
                name, subset, split = (spec.split("|") + ["", ""])[:3]
                st.phase(f"Streaming HF dataset {name}", total=args.max_rows or 0, unit="rows")
                try:
                    from datasets import load_dataset
                    ds = load_dataset(name, subset or None, split=split or "train", streaming=True)
                    n = 0
                    for rec in ds:
                        doc = classify_record(rec)
                        doc["src"] = f"hf:{name}"
                        emit(ef, doc)
                        n += 1
                        st.tick()
                        if n % 50000 == 0:
                            check_disk(out_dir, args.min_free_gb)
                            ef.flush()
                        if args.max_rows and n >= args.max_rows:
                            break
                    st.source("HF dataset rows", n)
                    print(f"  {name}: {n:,} rows", flush=True)
                    finish_source(f"hf:{spec}")
                except Exception as e:
                    st.warn(f"HF {name}: {e}")
                    st.drop("HF dataset failed to load")

            # ---------------- Wikipedia
            for line in args.wikipedia or []:
                if f"wiki:{line}" in done:
                    continue
                st.phase(f"Wikipedia: {line}", total=args.wiki_per_query if ":" in line else 1, unit="articles")
                try:
                    from dataset_sources import wikipedia_docs
                    for art in wikipedia_docs(line, args.wiki_per_query, args.wiki_lang):
                        emit(ef, {"kind": "text", "text": art["text"], "src": art["url"]})
                        st.source("Wikipedia article")
                        st.tick()
                    finish_source(f"wiki:{line}")
                except Exception as e:
                    st.warn(f"Wikipedia '{line}': {e}")
                    st.drop("Wikipedia request failed")

            # ---------------- ArXiv & Research Papers
            for line in getattr(args, "arxiv", []) or []:
                if f"arxiv:{line}" in done:
                    continue
                st.phase(f"Research papers: {line}", total=args.wiki_per_query, unit="papers")
                try:
                    from dataset_sources import arxiv_docs
                    count = 0
                    for paper in arxiv_docs(line, max_papers=args.wiki_per_query):
                        emit(ef, {"kind": "text", "text": paper["text"], "src": paper["url"]})
                        st.source("Research paper")
                        st.tick()
                        count += 1
                    if count == 0:
                        st.warn(f"No research papers found for '{line}'")
                    finish_source(f"arxiv:{line}")
                except Exception as e:
                    st.warn(f"Research query '{line}': {e}")
                    st.drop("Research query failed")

            # ---------------- web
            if args.url and "web" not in done:
                st.phase("Fetching web pages", total=len(args.url), unit="pages")
                for doc in fetch_urls(args.url, args.crawl_depth, args.max_pages, not args.allow_offsite, st, tmp_root,
                                      args.ocr, sitemap=getattr(args, "sitemap", False)):
                    emit(ef, doc)
                finish_source("web")

            # ---------------- YouTube + local audio/video (GPU Whisper, main process)
            transcriber = None

            def get_transcriber():
                nonlocal transcriber
                if transcriber is None:
                    from dataset_sources import Transcriber
                    st.phase(f"Loading Whisper ({args.whisper_model})")
                    transcriber = Transcriber(args.whisper_model)
                return transcriber

            for line in args.youtube or []:
                try:
                    from dataset_sources import youtube_videos, youtube_transcript, youtube_audio
                    videos = youtube_videos(line, args.max_videos)
                except Exception as e:
                    st.warn(f"YouTube '{line}': {e}")
                    st.drop("YouTube URL failed")
                    continue
                st.phase(f"YouTube transcripts ({len(videos)} videos)", total=len(videos), unit="videos")
                for vid, title in videos:
                    if f"yt:{vid}" in done:
                        st.tick()
                        continue
                    text = None
                    try:
                        text = youtube_transcript(vid)
                        st.source("YouTube transcript")
                    except Exception:
                        if args.transcribe:
                            try:
                                audio = youtube_audio(vid, os.path.join(tmp_root, "yt_audio"))
                                text = get_transcriber()(audio)
                                st.source("YouTube (Whisper)")
                                os.remove(audio)
                            except Exception as e:
                                st.warn(f"YouTube {vid}: {e}")
                        else:
                            st.drop("YouTube: no transcript (enable transcription)")
                    if text:
                        emit(ef, {"kind": "text", "text": f"{title}\n\n{text}", "src": f"https://youtu.be/{vid}"})
                    finish_source(f"yt:{vid}")
                    st.tick()

            todo_media = [m for m in media if f"media:{os.path.abspath(m)}" not in done]
            if todo_media:
                tr = get_transcriber()
                st.phase("Transcribing audio/video with Whisper", total=len(todo_media), unit="files")
                for m in todo_media:
                    try:
                        text = tr(m)
                        if text.strip():
                            emit(ef, {"kind": "text", "text": text, "src": m})
                            st.source(f"file {os.path.splitext(m)[1].lower()} (transcribed)")
                        else:
                            st.drop("no speech found")
                    except Exception as e:
                        st.warn(f"{os.path.basename(m)}: {e}")
                        st.drop("transcription failed")
                    finish_source(f"media:{os.path.abspath(m)}")
                    st.tick()
            if transcriber is not None:
                del transcriber
                _free_gpu()
        mark_done("extract")
    else:
        print("  extraction already complete - skipping", flush=True)

    # Recount kinds from disk (also covers resumed builds).
    kinds = Counter()
    with open(extracted_path, "rb") as ef:
        for raw in ef:
            m = re.match(rb'\{"kind": "(\w+)"', raw)
            if m:
                kinds[m.group(1).decode()] += 1
    if not sum(kinds.values()):
        raise RuntimeError("Nothing was extracted. Check your sources (see console for per-source errors).")

    # ================================================================== decide target format
    fmt = args.format
    if fmt == "auto":
        if kinds["sft"] and kinds["sft"] >= kinds["dpo"]:
            fmt = "sft"
        elif kinds["dpo"]:
            fmt = "dpo"
        else:
            fmt = "pretrain"
        if fmt == "sft" and kinds["text"] and not args.synthetic:
            st.warn(f"{kinds['text']:,} raw-text docs can't become chat data without Synthetic Q&A; they'll be skipped.")
    st.s["format"] = fmt
    print(f"  extracted {dict(kinds)} -> target format: {fmt}", flush=True)

    # ================================================================== 2. qualify (parallel)
    kept_path = os.path.join(build_dir, "kept.jsonl")
    pending_path = os.path.join(build_dir, "pending_chunks.jsonl")
    if not phase_done("qualify"):
        need_chunker = fmt == "pretrain" or (fmt == "sft" and args.synthetic)
        cfg = {"fmt": fmt, "synthetic": args.synthetic, "redact_pii": args.redact_pii, "min_chars": args.min_chars,
               "min_response_chars": args.min_response_chars, "system_prompt": args.system_prompt,
               "tokenizer": args.tokenizer, "overlap_tokens": args.overlap_tokens,
               "chunk_tokens": (args.chunk_tokens if fmt == "pretrain" else args.synth_chunk_tokens) if need_chunker else 0}
        dedup = Deduper(near=args.near_dedup, near_max=args.near_dedup_max, status=st)
        total = sum(kinds.values())
        st.phase("Cleaning, filtering & deduplicating", total=total, unit="docs")
        st.s["counters"]["examples kept"] = 0
        st.s["dropped"] = {k: v for k, v in st.s["dropped"].items()
                           if not (k.startswith("chunk") or k in _QUALIFY_REASONS)}  # don't double count on resume
        kept = 0
        workers = 1 if total < 3000 else args.workers
        pool = mp.get_context("spawn").Pool(workers, initializer=_qinit, initargs=(cfg,)) if workers > 1 else None
        if pool is None:
            _qinit(cfg)
        try:
            with open(extracted_path, encoding="utf-8") as ef, open(kept_path, "w", encoding="utf-8") as kf, \
                    open(pending_path, "w", encoding="utf-8") as pf:
                stop = False
                while not stop:
                    batch = [l for _, l in zip(range(4096), ef)]
                    if not batch:
                        break
                    results = pool.map(qualify_doc, batch, chunksize=32) if pool else [qualify_doc(l) for l in batch]
                    for res in results:
                        for item in res:
                            if item[0] == "drop":
                                st.drop(item[1])
                                continue
                            reason = dedup.is_dup(item[2])
                            if reason:
                                st.drop(reason if item[0] != "synth" and "text" not in item[1] else f"chunk: {reason}")
                                continue
                            if item[0] == "keep":
                                kf.write(json.dumps(item[1], ensure_ascii=False) + "\n")
                                kept += 1
                            else:
                                pf.write(json.dumps(item[1], ensure_ascii=False) + "\n")
                                st.count("chunks for synthesis")
                        if args.max_examples and kept >= args.max_examples:
                            stop = True
                            break
                    st.s["counters"]["examples kept"] = kept
                    st.tick(len(batch))
                    check_disk(out_dir, args.min_free_gb)
        finally:
            if pool:
                pool.terminate()
        mark_done("qualify")
    else:
        print("  qualification already complete - skipping", flush=True)

    # ================================================================== 3. synthetic Q&A (resumable)
    n_pending = sum(1 for _ in open(pending_path, "rb")) if os.path.exists(pending_path) else 0
    if n_pending and args.synthetic and not phase_done("synth"):
        budget = min(n_pending, args.max_synth_chunks or n_pending)
        if budget < n_pending:
            st.drop("chunk: over synthetic budget", n_pending - budget)
        # Deterministic sample of chunks so a resumed run picks the same ones.
        chosen = set(random.Random(args.seed).sample(range(n_pending), budget))
        raw_path = os.path.join(build_dir, "synthetic_raw.jsonl")
        already = sum(1 for _ in open(raw_path, "rb")) if os.path.exists(raw_path) else 0
        chunks = [json.loads(l) for i, l in enumerate(open(pending_path, encoding="utf-8")) if i in chosen][already:]
        st.phase(f"Loading generator {args.gen_model}")
        gen = QAGenerator(args.gen_model, load_4bit=args.gen_4bit, batch_size=args.gen_batch)
        st.phase("Generating Q&A pairs with local LLM", total=budget, unit="chunks")
        st.tick(already)
        q_seen = set()
        kept = sum(1 for _ in open(kept_path, "rb"))
        with open(kept_path, "a", encoding="utf-8") as kf, open(raw_path, "a", encoding="utf-8") as rawf:
            for i in range(0, len(chunks), gen.batch_size):
                batch = chunks[i:i + gen.batch_size]
                outs = gen.generate(batch, args.qa_per_chunk)
                for chunk, out in zip(batch, outs):
                    pairs = parse_qa(out)
                    if not pairs:
                        st.drop("synthetic: unparseable output")
                    for q, a in pairs:
                        qn = re.sub(r"\W+", " ", q.lower()).strip()
                        if len(q) < 10 or len(a) < args.min_response_chars:
                            st.drop("synthetic: too short")
                        elif BAD_Q.search(q) or BAD_Q.search(a[:80]):
                            st.drop("synthetic: references 'the passage'")
                        elif not dq.grounded(q, a, chunk):
                            st.drop("synthetic: answer not supported by the passage")
                        elif _auto_drop({"messages": [{"role": "user", "content": q}, {"role": "assistant", "content": a}]}):
                            st.drop("synthetic: failed quality checks")
                        elif qn in q_seen:
                            st.drop("synthetic: duplicate question")
                        else:
                            q_seen.add(qn)
                            user = f"{chunk}\n\n{q}" if args.qa_with_context else q
                            msgs = ([{"role": "system", "content": args.system_prompt}] if args.system_prompt else []) + \
                                   [{"role": "user", "content": user}, {"role": "assistant", "content": a}]
                            kf.write(json.dumps({"messages": msgs}, ensure_ascii=False) + "\n")
                            kept += 1
                            st.count("synthetic pairs")
                            st.s["counters"]["examples kept"] = kept
                    kf.flush()
                    rawf.write(json.dumps({"chunk": chunk, "output": out}, ensure_ascii=False) + "\n")
                    rawf.flush()
                    st.tick()
                print(f"  {already + min(i + gen.batch_size, len(chunks))}/{budget} chunks -> {kept:,} examples", flush=True)
        del gen
        _free_gpu()
        mark_done("synth")

    if not os.path.exists(kept_path) or os.path.getsize(kept_path) == 0:
        raise RuntimeError("Every record was filtered out - see the 'dropped' breakdown. "
                           "Try a lower --min_chars, a different --format, or --synthetic for raw documents.")

    # ================================================================== 4. finalize
    stats = finalize(kept_path, out_dir, build_dir, st, fmt, name=args.name, tokenizer_id=args.tokenizer,
                     eval_ratio=args.eval_ratio, seed=args.seed, output_format=args.output_format,
                     shard_rows=args.shard_rows, min_free_gb=args.min_free_gb,
                     extra={"extracted": dict(kinds), "config": st.s["config"], "warnings": st.s["warnings"]})
    if not args.keep_intermediate:
        shutil.rmtree(tmp_root, ignore_errors=True)
        for p in (extracted_path, kept_path, pending_path):
            if os.path.exists(p):
                os.remove(p)

    trainer = {"sft": "train_llm_sft.py", "dpo": "train_llm_dpo.py", "pretrain": "pretrain_llm.py"}[fmt]
    print(f"\n==> Done: {stats['train_examples']:,} train / {stats['eval_examples']:,} eval examples, "
          f"{stats['total_tokens']:,} tokens{' (est.)' if stats['tokens_estimated'] else ''}", flush=True)
    print(f"    Train with: python {trainer} --dataset {stats['train_path']}", flush=True)
    st.s.update(state="done", phase="Finished", stats=stats)
    st.write(True)


_QUALIFY_REASONS = {"too short", "mostly symbols/numbers", "repetitive lines", "repetitive words", "gibberish (word length)",
                    "web boilerplate", "exact duplicate", "near duplicate", "response too short", "refusal / AI boilerplate",
                    "conversation missing user or assistant turn", "invalid preference pair",
                    "raw text (enable synthetic Q&A to use)", "raw text (not used for DPO)",
                    "preference row (not used in this format)", "instruction row (not used in this format)"}


def _clean_dir(out_dir):
    """Wipe a previous build, but keep build/console.log (the Studio may be writing our stdout into it)."""
    keep = os.path.abspath(os.path.join(out_dir, "build", "console.log"))
    for root, dirs, files in os.walk(out_dir, topdown=False):
        for f in files:
            p = os.path.abspath(os.path.join(root, f))
            if p != keep:
                os.remove(p)
        for d in dirs:
            p = os.path.join(root, d)
            if not os.listdir(p):
                os.rmdir(p)


def _extract_job(path, ocr):
    docs, err = extract_document(path, ocr)
    return path, docs, err


def _free_gpu():
    import gc
    gc.collect()
    try:
        import torch
        torch.cuda.empty_cache()
    except Exception:
        pass


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="Build a training-ready LLM dataset from any sources",
                                formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    p.add_argument("--name", required=True, help="Output goes to data/<name>/")
    g = p.add_argument_group("sources")
    g.add_argument("--source", action="append", help="File, folder, glob or .zip (repeatable)")
    g.add_argument("--url", action="append", help="Web page / file URL (repeatable)")
    g.add_argument("--hf", action="append", help="HF dataset 'owner/name' or 'owner/name|subset|split' (repeatable)")
    g.add_argument("--github", action="append", help="GitHub repo 'owner/repo' or URL (repeatable)")
    g.add_argument("--wikipedia", action="append", help="'Title', 'search: query', 'category: Name' or URL (repeatable)")
    g.add_argument("--arxiv", action="append", help="ArXiv / research paper search query or ID (repeatable)")
    g.add_argument("--youtube", action="append", help="YouTube video / playlist / channel URL (repeatable)")
    g.add_argument("--crawl_depth", type=int, default=0, help="Follow links this many hops from each URL")
    g.add_argument("--max_pages", type=int, default=100)
    g.add_argument("--sitemap", action="store_true", help="Auto-discover sitemap.xml for web URLs")
    g.add_argument("--allow_offsite", action="store_true", help="Let the crawler leave the starting domains")
    g.add_argument("--wiki_per_query", type=int, default=20)
    g.add_argument("--wiki_lang", default="en")
    g.add_argument("--max_videos", type=int, default=50, help="Per playlist/channel")
    g.add_argument("--github_max_mb", type=int, default=2000)
    g.add_argument("--ocr", action="store_true", help="OCR images and scanned PDFs (RapidOCR, CPU)")
    g.add_argument("--transcribe", action="store_true", help="Transcribe audio/video (and YouTube w/o captions) with Whisper")
    g.add_argument("--whisper_model", default="openai/whisper-large-v3-turbo")
    q = p.add_argument_group("qualification")
    q.add_argument("--format", choices=["auto", "pretrain", "sft", "dpo"], default="auto")
    q.add_argument("--max_rows", type=int, default=0, help="Max rows per table file / HF dataset (0 = all)")
    q.add_argument("--max_examples", type=int, default=0, help="Stop after this many qualified examples (0 = all)")
    q.add_argument("--tokenizer", default="Qwen/Qwen2.5-0.5B-Instruct", help="Tokenizer for chunking & token stats")
    q.add_argument("--chunk_tokens", type=int, default=1024, help="Pretrain chunk size")
    q.add_argument("--overlap_tokens", type=int, default=64)
    q.add_argument("--min_chars", type=int, default=80)
    q.add_argument("--min_response_chars", type=int, default=20)
    q.add_argument("--no_near_dedup", dest="near_dedup", action="store_false")
    q.add_argument("--near_dedup_max", type=int, default=3_000_000, help="Disable near-dedup beyond this many docs (RAM)")
    q.add_argument("--redact_pii", action="store_true", help="Mask emails, phones, card numbers, IPs, API keys")
    q.add_argument("--system_prompt", default=None, help="Add this system message to every chat example")
    q.add_argument("--eval_ratio", type=float, default=0.02)
    s = p.add_argument_group("synthetic Q&A")
    s.add_argument("--synthetic", action="store_true", help="Turn raw documents into Q&A chat data with a local LLM")
    s.add_argument("--gen_model", default="Qwen/Qwen2.5-3B-Instruct")
    s.add_argument("--gen_4bit", action="store_true", help="Load the generator in 4-bit (for 7B+)")
    s.add_argument("--gen_batch", type=int, default=8)
    s.add_argument("--qa_per_chunk", type=int, default=3)
    s.add_argument("--synth_chunk_tokens", type=int, default=600)
    s.add_argument("--max_synth_chunks", type=int, default=2000, help="Cap on chunks sent to the generator (0 = all)")
    s.add_argument("--qa_with_context", action="store_true", help="Include the passage in the user turn (RAG-style)")
    o = p.add_argument_group("scale & output")
    o.add_argument("--output_format", choices=["jsonl", "parquet"], default="jsonl", help="parquet = zstd-compressed shards")
    o.add_argument("--shard_rows", type=int, default=500_000, help="Rows per parquet shard")
    o.add_argument("--resume", action="store_true", help="Continue an interrupted build with the same --name")
    o.add_argument("--overwrite", action="store_true", help="Replace an existing dataset with the same --name")
    o.add_argument("--min_free_gb", type=float, default=2.0, help="Stop safely if free disk drops below this")
    o.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 4) - 2))
    o.add_argument("--keep_intermediate", action="store_true")
    o.add_argument("--seed", type=int, default=42)
    a = p.parse_args(argv)
    if not (a.source or a.url or a.hf or a.github or a.wikipedia or a.youtube or a.arxiv or a.resume):
        p.error("give at least one source: --source, --url, --hf, --github, --wikipedia, --arxiv or --youtube")
    return a


if __name__ == "__main__":
    build(parse_args())
